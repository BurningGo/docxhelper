from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.enum.text import WD_COLOR_INDEX

def AlignAllCellsVerticallyCenter(table):
    n_row = len(table.rows)
    n_col = len(table.columns)

    for r in range(n_row):
        for c in range(n_col):
            tc = table.cell(r,c)._tc
            tcPr = tc.get_or_add_tcPr()
            tcVAlign = OxmlElement('w:vAlign')
            tcVAlign.set(qn('w:val'), "center")
            tcPr.append(tcVAlign)

def SetTableWidth_TwoColumns(table_obj, column0_width, column1_width):
    num_of_col = table_obj._column_count
    num_of_row = int(len(table_obj._cells) / num_of_col)

    width0 = Inches(column0_width)
    width1 = Inches(column1_width)

    for i in range(num_of_row):
        table_obj.cell(i,0).width = width0
        table_obj.cell(i,1).width = width1
    return table_obj


